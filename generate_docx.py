import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def clean_inline_formatting(text):
    # Clean LaTeX math markers for clean Word text
    text = text.replace(r'\text{', '').replace(r'}', '')
    text = text.replace(r'\ge', '≥').replace(r'\le', '≤')
    text = text.replace(r'\rightarrow', '→').replace(r'\longrightarrow', '→')
    text = text.replace(r'\Delta', 'Δ').replace(r'\lambda', 'λ')
    text = text.replace(r'\times', '×').replace(r'\pm', '±')
    text = text.replace(r'$$', '').replace(r'$', '')
    return text

def add_styled_paragraph(doc, text, style='Normal', space_after=6, space_before=0, bold=False, italic=False, font_size=10.5, color=RGBColor(30, 41, 59)):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    
    text = clean_inline_formatting(text)
    
    # Simple markdown bold parser **bold**
    tokens = re.split(r'(\*\*.*?\*\*)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
        else:
            # check for *italic*
            it_tokens = re.split(r'(\*.*?\*)', token)
            for it in it_tokens:
                if it.startswith('*') and it.endswith('*') and len(it) > 2:
                    run = p.add_run(it[1:-1])
                    run.italic = True
                    run.font.name = 'Calibri'
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = color
                else:
                    run = p.add_run(it)
                    run.bold = bold
                    run.italic = italic
                    run.font.name = 'Calibri'
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = color
    return p

def build_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    
    # Page setup (Letter, 0.8 inch margins)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    i = 0
    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i].strip()
        
        # Check if table
        if line.startswith('|') and line.endswith('|'):
            table_lines.append(line)
            i += 1
            continue
        elif table_lines:
            # Process accumulated table
            rows_data = []
            for t_line in table_lines:
                # ignore separator row like |:---|:---|
                if re.match(r'^\|[\s\-:]+(\|[\s\-:]+)+\|$', t_line):
                    continue
                cells = [c.strip() for c in t_line.strip('|').split('|')]
                rows_data.append(cells)
            
            if rows_data:
                col_count = max(len(r) for r in rows_data)
                table = doc.add_table(rows=len(rows_data), cols=col_count)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True
                
                # Style table
                for r_idx, row in enumerate(rows_data):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx < col_count:
                            cell = table.cell(r_idx, c_idx)
                            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                            cell_text = clean_inline_formatting(cell_text)
                            
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.line_spacing = 1.05
                            
                            # Parse bold in cell
                            tokens = re.split(r'(\*\*.*?\*\*)', cell_text)
                            for token in tokens:
                                if token.startswith('**') and token.endswith('**'):
                                    run = p.add_run(token[2:-2])
                                    run.bold = True
                                    run.font.name = 'Calibri'
                                    run.font.size = Pt(9.5)
                                else:
                                    run = p.add_run(token)
                                    run.font.name = 'Calibri'
                                    run.font.size = Pt(9.5)
                            
                            if r_idx == 0:
                                set_cell_background(cell, '1E293B')  # Slate 800
                                for r in p.runs:
                                    r.font.bold = True
                                    r.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, 'F8FAFC')  # Slate 50
                                else:
                                    set_cell_background(cell, 'FFFFFF')
                                for r in p.runs:
                                    r.font.color.rgb = RGBColor(30, 41, 59)
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            table_lines = []
            
        if not line:
            i += 1
            continue
            
        # Title (# )
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[2:])
            run.font.name = 'Arial'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)  # Slate 900
            
        # Heading 2 (## )
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[3:])
            run.font.name = 'Arial'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 76, 129)  # Deep Sapphire Blue
            
        # Heading 3 (### )
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)  # Indigo
            
        # Heading 4 (#### )
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)  # Slate 600
            
        # Divider (---)
        elif line == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('━' * 55)
            run.font.color.rgb = RGBColor(226, 232, 240)
            run.font.size = Pt(9)
            
        # Bullet list (* or - )
        elif line.startswith('* ') or line.startswith('- '):
            text = line[2:]
            p = add_styled_paragraph(doc, text, style='List Bullet', space_after=3, space_before=1, font_size=10.5)
            p.paragraph_format.left_indent = Inches(0.25)
            
        # Nested Bullet list (  - or   *)
        elif line.startswith('  - ') or line.startswith('  * '):
            text = line[4:]
            p = add_styled_paragraph(doc, text, style='List Bullet 2', space_after=2, space_before=1, font_size=10)
            p.paragraph_format.left_indent = Inches(0.5)
            
        # Numbered list (1. 2. etc)
        elif re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line)
            p = add_styled_paragraph(doc, text, style='List Number', space_after=3, space_before=1, font_size=10.5)
            p.paragraph_format.left_indent = Inches(0.25)
            
        # Regular text
        else:
            add_styled_paragraph(doc, line, space_after=5, space_before=1, font_size=10.5)
            
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated DOCX at: {docx_path}")

if __name__ == '__main__':
    md_file = r'c:\Users\ACT\Downloads\USI\PRODUCT_THINKING_DOCUMENT.md'
    docx_file = r'c:\Users\ACT\Downloads\USI\PRODUCT_THINKING_DOCUMENT.docx'
    build_docx(md_file, docx_file)
